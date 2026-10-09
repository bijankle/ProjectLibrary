"""Excel copies of the electrical lists for Sources (issues.html#sources), with their tables for the in-app viewer, the
same as tools/build_sources.py makes for the other lists: sources/<key>.xlsx and sources/<key>/<sheet>-0.json, entries in
sources/index.json and tools/doc_revs.json (and issues.json meta.docs, which mirrors doc_revs.json).
Each list's rows (tools/elec_rows.json) are grouped into sheets by kind: Cable schedule, Equipment list, Load list,
MCC schedule; columns as printed, plus the page each row is on.   python3 tools/build_elec_sources.py
"""
import json, os, re, shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
R = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
E = json.load(open(os.path.join(R, "tools", "elec_rows.json"))); DOCS = E["docs"]
SIX = os.path.join(R, "sources", "index.json"); six = json.load(open(SIX))
REVF = os.path.join(R, "tools", "doc_revs.json"); revs = json.load(open(REVF))
for k in [k for k in six if k.startswith("el")]: six.pop(k)
for k in [k for k in revs if k.startswith("el") and revs[k].get("elec")]: revs.pop(k)
def kind(h):
    s = " ".join(h).lower()
    if "cable number" in s or "cable no" in s: return "Cable schedule"
    if "module" in s or "circuit breaker" in s: return "MCC schedule"
    if "load type" in s or "nameplate (kw)" in s: return "Load list"
    return "Equipment list"
def nice(t):
    t = re.sub(r"\s+", " ", t).strip()
    return t[:1] + t[1:].lower() if t.isupper() else t
byd = {}
for x in E["rows"]: byd.setdefault(x["doc"], []).append(x)
CH = 5000
for doc, rows in byd.items():
    m = re.match(r"2000-([A-Z0-9]+)-LST-EL-(\d+)", doc); key = f"el{m.group(1).lower()}{m.group(2)}"
    sheets = {}
    for x in rows:
        k = kind(list(x["r"].keys())); sh = sheets.setdefault(k, {"cols": [], "rows": []})
        for h in x["r"]:
            if h not in sh["cols"]: sh["cols"].append(h)
        sh["rows"].append(x)
    d = DOCS[doc]; title = nice(d["title"]); fname = f"{doc} {title} Rev {d['rev']} (tables).xlsx"
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    os.makedirs(os.path.join(R, "sources", key), exist_ok=True)
    for f in os.listdir(os.path.join(R, "sources", key)): os.remove(os.path.join(R, "sources", key, f))
    meta = []
    for i, (name, sh) in enumerate(sheets.items()):
        cols = sh["cols"] + ["Page"]
        data = [cols] + [[x["r"].get(c, "") for c in sh["cols"]] + [x["pg"]] for x in sh["rows"]]
        ws = wb.create_sheet(name[:31]); ws.freeze_panes = "A2"
        for r in data: ws.append(r)
        for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="FCE4B6"); c.alignment = Alignment(wrap_text=True, vertical="top")
        for j, c in enumerate(cols): ws.column_dimensions[openpyxl.utils.get_column_letter(j + 1)].width = min(40, max(10, len(str(c)) + 2))
        ws.auto_filter.ref = ws.dimensions
        json.dump(data, open(os.path.join(R, "sources", key, f"{i}-0.json"), "w"), ensure_ascii=False, separators=(",", ":"))
        meta.append({"name": name, "hdr": 0, "n": len(data), "c": len(cols), "chunks": 1})
    path = os.path.join(R, "sources", key + ".xlsx"); wb.save(path)
    six[key] = {"chunk": CH, "xlsx": f"sources/{key}.xlsx", "fname": fname, "how": "The tables of the PDF list, read row by row (tools/build_elec.py)", "from": "PDF",
                "sheets": meta, "size": os.path.getsize(path)}
    revs[key] = {"title": title, "full": title, "number": doc, "kind": "List", "rev": d["rev"], "date": d["date"], "status": "", "elec": 1,
                 "used": "Electrical equipment, cables and the electrical data of MEL equipment", "history": [{"rev": d["rev"], "date": d["date"], "status": ""}],
                 "note": "PDF list; tables read from the PDF.", "project": "TSF" if re.search(r"-EL-4\d{4}$|-T\d\d-", doc) else ""}
json.dump(six, open(SIX, "w"), ensure_ascii=False, indent=1)
json.dump(revs, open(REVF, "w"), ensure_ascii=False, indent=1)
I = json.load(open(os.path.join(R, "issues.json"))); I["meta"]["docs"] = revs
json.dump(I, open(os.path.join(R, "issues.json"), "w"), ensure_ascii=False, separators=(",", ":"))
print(len(byd), "electrical lists as Excel")
