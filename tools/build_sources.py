"""Builds the source document tables for the Source documents tab (issues.html#sources): a cleaned Excel copy of every
source document and a JSON copy of each of its tables for the in-app viewer.

Usage: python3 tools/build_sources.py <folder holding the source .xlsx and .pdf files>   (run tools/doc_revs.py first)
Output: sources/index.json, sources/<key>.xlsx, sources/<key>/<sheet>-<chunk>.json (5,000 rows per chunk)

Spreadsheets: every sheet except the cover page is rebuilt from its cell values (as last calculated), so all table data
and columns are kept (hidden sheets stay hidden), and the cover page, logos, images, headers and footers, comments,
file properties, external links and data connections are left behind. People's names found on the cover pages or in
the file properties are also blanked if they appear anywhere in the tables.
PDFs: the data tables are pulled out into one sheet each. Report cover pages and the revision / sign off tables are
skipped. The PFDs are drawings, not tables, so three tables are built from their text: sheets, equipment (tag and name
under each symbol) and off sheet connections. The overall plant layout gives its area legend.
"""
import datetime, glob, json, os, re, shutil, sys, warnings
from collections import defaultdict
import openpyxl
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
warnings.filterwarnings("ignore")
SRC = sys.argv[1]
HERE = os.path.dirname(__file__)
OUT = os.path.join(HERE, "..", "sources")
REV = json.load(open(os.path.join(HERE, "doc_revs.json")))
FILEKEY = {"mel": "MEL_2000-F00-LST-ME-10001", "ins": "2000-F00-LST-IC-11011", "cv": "2000-F00-LST-PP-10003", "line": "2000-F00-LST-PP-10001",
           "mv": "2000-F00-LST-PP-10002", "spi": "2000-F00-LST-PP-10005", "hose1": "2000-F00-LST-PP-10032", "hose2": "2000-F00-LST-PP-12032",
           "pdc": "2000-F00-DCR-PR-10002", "cp1": "2000-F00-REP-PR-10001", "cp2": "2000-F00-REP-PR-10002", "pfd": "PFDs_combined",
           "num": "2000-F00-STS-GE-10003", "layout": "2000-F00-DRG-GE-20001", "wbs": "KCGM_Growth_WBS"}
SIGNOFF = re.compile(r"^(PREPARED|REVIEWED|APPROVED|DRAWN|DRAW|CHECKED|CHECK|DESIGN|TECH APP|PROJ APP|SIGNED|SIGNATURE)$", re.I)

def find(k):
    m = [f for f in glob.glob(os.path.join(SRC, "*")) if FILEKEY[k] in os.path.basename(f) and f.endswith((".xlsx", ".pdf"))]
    return m[0] if m else None

# ---------- names to keep out ----------
NAMES = set()
def add_name(s):
    s = re.sub(r"\s*\(.*?\)", "", str(s or "")).strip()
    if re.match(r"^[A-Z][A-Za-z'\-]*\.? [A-Z][A-Za-z'\-]+( [A-Z][A-Za-z'\-]+)?$", s):
        NAMES.add(s); p = s.split(); NAMES.add(p[0][0] + " " + p[-1])        # "Luke Visaggio" and "L Visaggio"
for k in FILEKEY:
    f = find(k)
    if not f: continue
    if f.endswith(".xlsx"):
        import zipfile
        core = zipfile.ZipFile(f).read("docProps/core.xml").decode("utf8", "ignore")
        for n in re.findall(r"<(?:dc:creator|cp:lastModifiedBy)>([^<]+)<", core): add_name(n)
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        for s in wb.sheetnames:
            if "cover" not in s.lower(): continue
            hdr = None
            for r in wb[s].iter_rows(values_only=True):
                v = [x for x in r if x not in (None, "")]
                if [str(x).upper() for x in v[:6]] == ["DATE", "REVISION", "STATUS", "PREPARED", "REVIEWED", "APPROVED"]: hdr = True; continue
                if hdr and len(v) >= 6: [add_name(x) for x in v[3:6]]
    else:
        import pymupdf
        t = pymupdf.open(f)[0].get_text()
        for d, r, st, a, b, c in re.findall(r"\n(\d{1,2}-[A-Za-z]{3,4}-\d{2})\s*\n(\w{1,2})\s*\n([^\n]+)\n(?:[^\n]*\n)??([A-Z]\.? ?[A-Za-z]+ [A-Za-z'\-]+)\s*\n([A-Z]\.? ?[A-Za-z]+ [A-Za-z'\-]+)\s*\n([A-Z]\.? ?[A-Za-z]+ [A-Za-z'\-]+)", t):
            [add_name(x) for x in (a, b, c)]
        for n in re.findall(r"\n([A-Z] [A-Z][a-z]+)\s*\n\d{1,2}/\d{1,2}/\d{2,4}", t): add_name(n)   # client review stamp signatory
NAMES.discard("")
NAME_RE = re.compile(r"\b(" + "|".join(sorted(map(re.escape, NAMES), key=len, reverse=True)) + r")\b", re.I) if NAMES else None
scrubbed = defaultdict(int)

def val(v, key):
    """Cell value for the JSON viewer and the clean workbook: text, number or ISO date; names blanked."""
    if v is None: return ""
    if isinstance(v, datetime.datetime): return v.date().isoformat() if v.time() == datetime.time() else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, (datetime.date, datetime.time)): return v.isoformat()
    if isinstance(v, bool): return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")): return ""
        return int(v) if v == int(v) and abs(v) < 1e15 else round(v, 6)
    if isinstance(v, int): return v
    s = str(v)
    if s.startswith("#") and s.rstrip("!?/0").upper() in ("#N/A", "#REF", "#VALUE", "#DIV", "#NAME", "#NUM", "#NULL"): return s
    if NAME_RE and NAME_RE.search(s): scrubbed[key] += len(NAME_RE.findall(s)); s = NAME_RE.sub("[name removed]", s)
    return s

def trim(rows):
    while rows and not any(x != "" for x in rows[-1]): rows.pop()
    w = max((max((i + 1 for i, x in enumerate(r) if x != ""), default=0) for r in rows), default=0)
    return [list(r[:w]) + [""] * (w - len(r[:w])) for r in rows]

def header_row(rows):
    best, bi = 0, 0
    for i, r in enumerate(rows[:15]):
        n = sum(1 for x in r if x != "")
        if n > best: best, bi = n, i
    return bi

HFONT, HFILL = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="3A3226")
def write_book(path, sheets, title, note):
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("About")
    for r in ([title], [note], [], ["Sheet", "Rows", "Columns", "Hidden in the original"]) + tuple([s["name"], len(s["rows"]), len(s["rows"][0]) if s["rows"] else 0, "Yes" if s.get("hidden") else ""] for s in sheets):
        ws.append(r)
    for s in sheets:
        rows, h = s["rows"], s["hdr"]
        ws = wb.create_sheet(s["name"][:31])
        if s.get("hidden"): ws.sheet_state = "hidden"
        widths = [8] * (len(rows[0]) if rows else 0)
        for r in rows[:400]:
            for i, x in enumerate(r): widths[i] = min(60, max(widths[i], len(str(x)) * 1.1 + 2 if len(str(x)) < 80 else 60))
        for i, w in enumerate(widths): ws.column_dimensions[get_column_letter(i + 1)].width = w
        if rows: ws.freeze_panes = f"A{h + 2}"
        for ri, r in enumerate(rows):
            if ri == h:
                out = []
                for x in r:
                    c = WriteOnlyCell(ws, value=x); c.font, c.fill, c.alignment = HFONT, HFILL, Alignment(wrap_text=True, vertical="top"); out.append(c)
                ws.append(out)
            else: ws.append([x if x != "" else None for x in r])
        if rows and len(rows) > h + 1 and rows[0]: ws.auto_filter.ref = f"A{h + 1}:{get_column_letter(len(rows[0]))}{len(rows)}"
    wb.save(path)

# ---------- spreadsheets ----------
def xlsx_tables(k, f):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        if "cover" in ws.title.lower(): continue
        rows = trim([[val(v, k) for v in r] for r in ws.iter_rows(values_only=True)])
        if not rows: continue
        out.append(dict(name=ws.title, rows=rows, hdr=header_row(rows), hidden=ws.sheet_state != "visible"))
    return out

# ---------- PDFs ----------
def clean_table(t):
    rows = [[re.sub(r"\s*\n\s*", " ", str(x)).strip() if x is not None else "" for x in r] for r in t]
    rows = [r for r in rows if any(r)]
    keep = [i for i in range(max((len(r) for r in rows), default=0)) if any(i < len(r) and r[i] for r in rows)]
    return [[r[i] if i < len(r) else "" for i in keep] for r in rows]

def line_rows(page, t):
    """Tables drawn without row lines (the PDC data sheets) come out as one tall cell per column. Rebuild them one text
    line per row: words are grouped by line and dropped into the column whose left edge they sit right of."""
    xs = sorted({round(c[0]) for c in t.cells if c})
    edges = [x for i, x in enumerate(xs) if i == 0 or x - xs[i - 1] > 3]
    x0, y0, x1, y1 = t.bbox
    words = [w for w in page.get_text("words") if w[0] >= x0 - 1 and w[2] <= x1 + 1 and w[1] >= y0 - 1 and w[3] <= y1 + 1]
    words.sort(key=lambda w: ((w[1] + w[3]) / 2, w[0]))
    lines, cur, cy = [], [], None
    for w in words:
        y = (w[1] + w[3]) / 2
        if cy is not None and abs(y - cy) > 2.5: lines.append(cur); cur = []
        cur.append(w); cy = y if not cur[:-1] else cy
    if cur: lines.append(cur)
    rows = []
    for ln in lines:
        r = [""] * len(edges)
        for w in sorted(ln, key=lambda w: w[0]):
            ci = max(i for i, e in enumerate(edges) if w[0] >= e - 2) if w[0] >= edges[0] - 2 else 0
            r[ci] = (r[ci] + " " + w[4]).strip()
        rows.append(r)
    return rows

def pdf_tables(k, f):
    import pymupdf
    doc, out = pymupdf.open(f), []
    for pi, p in enumerate(doc):
        if pi == 0 and k != "wbs": continue      # cover page: title and sign off only
        for ti, t in enumerate(p.find_tables().tables):
            raw = t.extract()
            tall = any(str(x).count("\n") > 4 for r in raw for x in r if x)
            rows = clean_table(line_rows(p, t) if tall else raw)
            if len(rows) < 2 or not rows[0]: continue
            if any(SIGNOFF.match(x) for r in rows[:3] for x in r): continue
            rows = [[val(x, k) for x in r] for r in rows]
            out.append(dict(name=f"p{pi + 1} t{ti + 1}", rows=rows, hdr=0, page=pi + 1, caption=" | ".join(x for x in rows[0] if x)[:90]))
    return out

TAG = re.compile(r"^[FNTM]\d{2}-[A-Z]{1,4}-\d{2,4}[A-Z]?(/\w+)?$")
DRG = re.compile(r"^2000-F00-PFD-PR-\d{5}$")
def pfd_tables(k, f):
    import pymupdf
    sheets, equip, conns = [["Drawing number", "Title", "Rev", "Date", "Status"]], [["Drawing number", "Tag", "Name on the drawing"]], [["Drawing number", "Connects to drawing", "Destination / source", "Stream"]]
    for p in pymupdf.open(f):
        lines = []
        for b in p.get_text("dict")["blocks"]:
            for l in b.get("lines", []):
                s = " ".join(sp["text"] for sp in l["spans"]).strip()
                if s: lines.append((l["bbox"], s))
        H = p.rect.height
        body = [(bb, s) for bb, s in lines if bb[1] < H * 0.9 and bb[0] > 40 and not s.startswith(("O:\\", "C:\\"))]
        full = p.get_text().replace("\n", "|")
        nums = re.findall(r"N\.T\.S\.\|(2000-F00-PFD-PR-\d{5})", full) or re.findall(r"(2000-F00-PFD-PR-\d{5})\|1\|0\|", full)
        if not nums: continue
        me = nums[0]
        tm = re.search(r"PRIMERO PROJECT NUMBER:\|(.+?)\|PROCESS FLOW DIAGRAM", full)
        revs = re.findall(r"\|\s*([0-9A-Z])\|\s*(\d{2}[A-Z]{3}\d{2})\s*\|?\s*([A-Z0-9 ]+?)\|", full)
        last = max(revs, key=lambda r: r[0]) if revs else ("", "", "")
        sheets.append([me, re.sub(r"\s*\|\s*", " ", tm[1]).strip() if tm else "", last[0], last[1], last[2]])
        def near_below(bb, n=4):
            cx, out, y = (bb[0] + bb[2]) / 2, [], bb[3]
            for b2, s2 in sorted(body, key=lambda z: z[0][1]):
                if b2[1] >= y - 2 and b2[1] - y < 16 and abs((b2[0] + b2[2]) / 2 - cx) < 75 and not TAG.match(s2) and not DRG.match(s2):
                    out.append(s2); y = b2[3]
                    if len(out) >= n: break
            return out
        for bb, s in body:
            for tag in re.findall(r"\b([FNTM]\d{2}-[A-Z]{1,4}-\d{2,4}[A-Z]?)\b", s):
                if TAG.match(s.split(" ")[-1]) or TAG.match(s):
                    equip.append([me, tag, val(" ".join(near_below(bb)), k)])
            for d in re.findall(r"2000-F00-PFD-PR-\d{5}", s):
                if d == me: continue
                above = [s2 for b2, s2 in body if 0 < bb[1] - b2[3] < 8 and abs(b2[0] - bb[0]) < 60]
                conns.append([me, d, val(above[-1] if above else "", k), val(" ".join(near_below(bb, 2)), k)])
    seen, eq = set(), []
    for r in equip:
        if (r[0], r[1]) not in seen: seen.add((r[0], r[1])); eq.append(r)
    return [dict(name="Sheets", rows=sheets, hdr=0), dict(name="Equipment", rows=eq, hdr=0), dict(name="Off sheet connections", rows=conns, hdr=0)]

def layout_tables(k, f):
    import pymupdf
    t = pymupdf.open(f)[0].get_text().replace("\n", "|")
    rows = [["Item", "Area", "Area description"]] + [[int(a), b, val(c.strip(), k)] for a, b, c in re.findall(r"\|(\d{1,2})\|([FNT]\d{2})\|([A-Z][^|]{3,})(?=\|)", t)]
    return [dict(name="Area legend", rows=rows, hdr=0)]

# ---------- build ----------
CHUNK = 5000
if os.path.isdir(OUT): shutil.rmtree(OUT)
os.makedirs(OUT)
index = {}
for k, r in REV.items():
    f = find(k)
    if not f: print("missing", k); continue
    if f.endswith(".xlsx"): tabs, how = xlsx_tables(k, f), "Every sheet of the original except the cover page, rebuilt from the cell values"
    elif k == "pfd": tabs, how = pfd_tables(k, f), "Built from the drawing text: sheet list, equipment tags and names, off sheet connections"
    elif k == "layout": tabs, how = layout_tables(k, f), "The area legend of the drawing"
    else: tabs, how = pdf_tables(k, f), "Every data table in the PDF (cover page and sign off tables left out), one sheet each"
    if not tabs: print("no tables", k); continue
    names = set()
    for t in tabs:   # unique sheet names of up to 31 characters
        n = re.sub(r"[\[\]:*?/\\]", " ", t["name"])[:31].strip() or "Sheet"
        while n in names: n = n[:28] + f" {len(names)}"
        names.add(n); t["name"] = n
    title = f"{r['number']} {r['full']}" + (f" Rev {r['rev']}" if r["rev"] else "")
    note = f"Tables only, from the original {os.path.splitext(f)[1][1:].upper()} ({how}). Cover page, logos, sign off names and file metadata removed. Built by tools/build_sources.py."
    write_book(os.path.join(OUT, k + ".xlsx"), tabs, title, note)
    os.makedirs(os.path.join(OUT, k))
    sh = []
    for i, t in enumerate(tabs):   # viewer copy in chunks of CHUNK rows, so a phone only loads what it shows
        rows = t["rows"]; nc = max(1, -(-len(rows) // CHUNK))
        for c in range(nc):
            json.dump(rows[c * CHUNK:(c + 1) * CHUNK], open(os.path.join(OUT, k, f"{i}-{c}.json"), "w"), ensure_ascii=False, separators=(",", ":"))
        sh.append({k2: t[k2] for k2 in ("name", "hdr", "hidden", "page", "caption") if t.get(k2) not in (None, False, "")} | {"n": len(rows), "c": len(rows[0]) if rows else 0, "chunks": nc})
    fname = re.sub(r"[^\w .()&-]", "", f"{title} (tables).xlsx")
    index[k] = {"chunk": CHUNK, "xlsx": f"sources/{k}.xlsx", "fname": fname, "how": how, "from": os.path.splitext(f)[1][1:].upper(), "sheets": sh,
                "size": os.path.getsize(os.path.join(OUT, k + ".xlsx"))}
    print(f"{k:7} {len(tabs):3} tables  {sum(t['n'] for t in sh):7} rows  {index[k]['size'] / 1e6:5.2f} MB  names blanked {scrubbed[k]}")
json.dump(index, open(os.path.join(OUT, "index.json"), "w"), ensure_ascii=False, indent=1)
print("names looked for:", len(NAMES))
