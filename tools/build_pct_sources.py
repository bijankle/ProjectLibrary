"""The Process Control Tables (sources/pct/ in the private repository) in Sources (issues.html#sources): a cleaned Excel
copy of each and a JSON copy of each table for the in-app viewer, the same as tools/build_sources.py makes for the other
lists: sources/<key>.xlsx and sources/<key>/<sheet>-<chunk>.json, entries in sources/index.json and tools/doc_revs.json
(and issues.json meta.docs, which mirrors doc_revs.json).   python3 tools/build_pct_sources.py

Kept: the revision history and every control table (analogue and digital I/O, totalisers, valves, PID controllers,
drives, sequences, calculations, misc, alarm database, cause and effect, risk matrix, reference list). Left out: the cover
page, contents and progress pages, reviewers' comment sheets, scratch sheets, and the copies of the instrument and
equipment lists some workbooks carry (the app has those lists in their own right). Sign off columns are blanked.
"""
import json, os, re, glob, datetime, collections, warnings
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
warnings.filterwarnings("ignore")
R = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RP = os.environ.get("KCGM_PRIVATE") or os.path.join(os.path.dirname(R), "projectlibraryprivate")
SIX = os.path.join(R, "sources", "index.json"); six = json.load(open(SIX))
REVF = os.path.join(R, "tools", "doc_revs.json"); revs = json.load(open(REVF))
for k in [k for k in six if k.startswith("pct")]: six.pop(k)
for k in [k for k in revs if k.startswith("pct")]: revs.pop(k)
DROP = re.compile(r"cover|^toc$|progress|comments|^sheet\d*$|to do|do not print|broken|^ref |^inputs$", re.I)
CH = 5000
# titles from the project document register (as tools/build_refonly.py reads it)
REGF = max(glob.glob(os.path.join(RP, "sources", "register", "*.xls[xm]")), key=os.path.getmtime)
REG = {str(r[3]).strip(): re.sub(r"\s+", " ", str(r[6] or "")).strip() for r in list(openpyxl.load_workbook(REGF, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))[1:] if r[3]}

def val(v):
    if v is None: return ""
    if isinstance(v, datetime.datetime): return v.date().isoformat() if v.time() == datetime.time() else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, (datetime.date, datetime.time)): return v.isoformat()
    if isinstance(v, bool): return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")): return ""
        return int(v) if v == int(v) and abs(v) < 1e15 else round(v, 6)
    return v if isinstance(v, int) else str(v)

def trim(rows):
    while rows and not any(x != "" for x in rows[-1]): rows.pop()
    w = max((max((i + 1 for i, x in enumerate(r) if x != ""), default=0) for r in rows), default=0)
    return [list(r[:w]) + [""] * (w - len(r[:w])) for r in rows]

docs = {}
for f in glob.glob(os.path.join(RP, "sources", "pct", "*.xls[xm]")):
    m = re.match(r"(2000-F\d{2}-LST-PR-\d{5})_(\w+?)_", os.path.basename(f))
    if m and (m.group(1) not in docs or m.group(2) > docs[m.group(1)][0]): docs[m.group(1)] = (m.group(2), f)
for dn, (rev, f) in sorted(docs.items()):
    key = "pct" + re.sub(r"\D", "", dn.split("-", 1)[1])   # pct1310001
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    cov = [re.sub(r"\s+", " ", str(c)) for r in wb.worksheets[0].iter_rows(values_only=True, max_row=30) for c in r if c]
    title = REG.get(dn) or next((c for c in cov if c.upper().startswith("PROCESS CONTROL TABLES - ")), "PROCESS CONTROL TABLES")
    title = re.sub(r"\b(And|Of|For|The)\b", lambda m: m.group(0).lower(), title.title())
    title = re.sub(r"\b(Cil|Ufg|Ew|Fcs)(?=\d|\b)|\b[A-Za-z]+\d+\w*", lambda m: m.group(0).upper(), title)
    sheets, dates = [], []
    for ws in wb.worksheets:
        if DROP.search(ws.title.strip()) or ws.sheet_state != "visible" and not re.match(r"\d+\.", ws.title): continue
        rows = [[val(c) for c in r[:80]] for r in ws.iter_rows(values_only=True)]
        rows = trim(rows)
        if not rows or len(rows) < 2: continue
        # sign off columns (initials and dates under "C1 Sign Off" / "C2 Sign Off") blanked
        so = set()
        for r in rows[:8]:
            for i, x in enumerate(r):
                if re.match(r"C[12] Sign Off", str(x)): so |= {i, i + 1}
        if so: rows = [[("" if i in so and j > 5 else x) for i, x in enumerate(r)] for j, r in enumerate(rows)]
        if ws.title.lower().startswith("revision"):
            dates += [x for r in rows for x in r if isinstance(x, str) and re.fullmatch(r"\d{4}-\d\d-\d\d", x)]
        hdr = next((i for i, r in enumerate(rows[:12]) if any(str(x).strip() in ("Tag Name", "Seq No.", "Revision") for x in r)), 0)
        sheets.append({"name": ws.title.strip()[:31], "rows": rows, "hdr": hdr})
    os.makedirs(os.path.join(R, "sources", key), exist_ok=True)
    for x in os.listdir(os.path.join(R, "sources", key)): os.remove(os.path.join(R, "sources", key, x))
    out = openpyxl.Workbook(); out.remove(out.active); meta = []
    for i, s in enumerate(sheets):
        ws = out.create_sheet(s["name"])
        for r in s["rows"]: ws.append(r)
        for c in (ws[s["hdr"] + 1] if s["rows"] else []): c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="FCE4B6"); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = ws.cell(row=s["hdr"] + 2, column=1)
        rows = s["rows"]; n = (len(rows) + CH - 1) // CH
        for c in range(n): json.dump(rows[c * CH:(c + 1) * CH], open(os.path.join(R, "sources", key, f"{i}-{c}.json"), "w"), ensure_ascii=False, separators=(",", ":"))
        meta.append({"name": s["name"], "hdr": s["hdr"], "n": len(rows), "c": len(rows[0]), "chunks": n})
    path = os.path.join(R, "sources", key + ".xlsx"); out.save(path)
    date = max(dates) if dates else ""
    six[key] = {"chunk": CH, "xlsx": f"sources/{key}.xlsx", "fname": f"{dn} {title} Rev {rev} (tables).xlsx",
                "how": "Every control table of the original workbook rebuilt from the cell values (cover, contents, comment sheets and list copies left out; sign off columns blanked)",
                "from": "XLSX", "sheets": meta, "size": os.path.getsize(path)}
    revs[key] = {"title": "PCT " + dn.split("-")[1], "full": title, "number": dn, "kind": "List", "rev": rev, "date": date, "status": "", "pct": 1,
                 "used": "Control data on each asset's page: ranges, set points, alarms, interlocks, permissives, controllers",
                 "history": [{"rev": rev, "date": date, "status": ""}], "note": "Process Control Tables workbook (" + os.path.basename(f).split("_", 1)[1].rsplit(".", 1)[0] + ")", "project": ""}
    print(key, dn, rev, date, title, len(sheets), "sheets")
json.dump(six, open(SIX, "w"), ensure_ascii=False, indent=1)
json.dump(revs, open(REVF, "w"), ensure_ascii=False, indent=1)
I = json.load(open(os.path.join(R, "issues.json"))); I["meta"]["docs"] = revs
json.dump(I, open(os.path.join(R, "issues.json"), "w"), ensure_ascii=False, separators=(",", ":"))
