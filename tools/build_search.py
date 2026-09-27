"""Builds search-data.json (used by lookup.js) straight from the project spreadsheets.

Usage: python3 tools/build_search.py <folder holding the .xlsx files>
Files are found by document number in the file name, so revisions can be dropped in:
  MEL 2000-F00-LST-ME-10001, Instrument List 2000-F00-LST-IC-11011, Control Valve List 2000-F00-LST-PP-10003,
  Line List 2000-F00-LST-PP-10001, Manual Valve List 2000-F00-LST-PP-10002, SPI List 2000-F00-LST-PP-10005,
  Slurry Hose Lists 2000-F00-LST-PP-10032 (Stage 1) and 2000-F00-LST-PP-12032 (Stage 2).
"""
import glob, json, os, re, sys, warnings
import openpyxl
warnings.filterwarnings("ignore")
SRC = sys.argv[1]
OUT = os.path.join(os.path.dirname(__file__), "..", "search-data.json")

def find(doc):
    m = [f for f in glob.glob(os.path.join(SRC, "*.xlsx")) if doc in f]
    if not m: sys.exit("missing " + doc)
    return m[0]

def clean(v):
    if v is None: return ""
    if isinstance(v, float):
        v = round(v, 3)
        if v == int(v): v = int(v)
    s = re.sub(r"\s+", " ", str(v)).strip()
    return "" if s in ("-", "N/A", "n/a", "NA", "None", "nan", "#N/A", "TBA", "TBC", "?", "0 mm") else s

def rows(doc, sheet, hdr):
    ws = openpyxl.load_workbook(find(doc), read_only=True, data_only=True)[sheet]
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i > hdr: yield r

TYPES, DATA = {}, {}
def add_type(key, name, doc, rev, fields):
    TYPES[key] = {"n": name, "doc": doc, "rev": rev, "f": [f[0] for f in fields]}
    DATA[key] = []
def put(key, fields, r):
    rec = [clean(r[c]) if c < len(r) else "" for _, c in fields]
    while rec and rec[-1] == "": rec.pop()
    DATA[key].append(rec)

# ---- MEL (key field first in every type) ----
F = [("Equipment number", 3), ("Equipment name", 4), ("Area", 2), ("Size / description", 8), ("Nominal duty point", 6), ("Design duty point", 7),
     ("Material of construction", 5), ("Installed power (kW)", 18), ("Voltage", 16), ("Starter type", 17), ("Duty / standby", 19), ("Emergency power", 20),
     ("Status", 10), ("Stage", 12), ("Brownfield / greenfield", 9), ("Supply type", 11), ("Package", 13), ("Vendor", 14), ("Make / model", 15),
     ("PFD", 22), ("P&ID", 23), ("Reference drawing", 24), ("MEL revision", 1)]
add_type("mel", "Equipment (MEL)", "2000-F00-LST-ME-10001", "Rev 0", F)
seen = set()
for r in rows("2000-F00-LST-ME-10001", "MEL", 2):
    e = clean(r[3])
    if not e or e in seen: continue
    seen.add(e); put("mel", F, r)

# ---- Instruments ----
F = [("Instrument tag", 23), ("Description", 24), ("Instrument type", 26), ("Equipment number", 6), ("Equipment description", 7), ("Location", 8),
     ("Loop number", 12), ("Process fluid", 27), ("Process connection", 28), ("Range / units", 29), ("Make", 30), ("Model", 31), ("Burnout", 32),
     ("Supplied by", 33), ("Installed by", 34), ("Remote IO panel", 35), ("Emergency power", 36), ("Valve type", 15), ("Valve size", 17),
     ("Stage", 2), ("Package", 4), ("P&ID", 25), ("Comments", 37)]
add_type("ins", "Instrument", "2000-F00-LST-IC-11011", "", F)
seen = set()
for r in rows("2000-F00-LST-IC-11011", "MASTER Instrument List", 1):
    t = clean(r[23])
    if not t or t in seen or t.startswith("#"): continue
    seen.add(t); put("ins", F, r)

# ---- Control valves ----
F = [("Valve tag", 4), ("Location", 8), ("Valve type", 1), ("Valve code", 7), ("Valve size (mm)", 5), ("Valve class", 6), ("Line number", 9),
     ("Service fluid", 14), ("Fluid", 52), ("Line spec", 15), ("Line size (mm)", 16), ("Flow min (m³/h)", 17), ("Flow max (m³/h)", 18),
     ("Inlet pressure (kPag)", 19), ("ΔP at min flow (kPa)", 20), ("ΔP at max flow (kPa)", 21), ("Solids P50 (µm)", 22), ("Solids P80 (µm)", 23),
     ("% solids (w/w)", 24), ("Fluid SG", 25), ("Max pressure (kPag)", 26), ("Closing time (s)", 27), ("Fail position", 29), ("Actuator type", 30),
     ("Solenoid type", 31), ("Solenoid tag", 32), ("Positioner", 33), ("Position feedback", 34), ("Special requirements", 35),
     ("Pneumatic panel / HPU", 36), ("Valve make", 41), ("Valve model", 42), ("Valve drawing", 43), ("Actuator make", 44), ("Actuator model", 45),
     ("Limit switches / positioner", 46), ("Package", 38), ("Stage", 12), ("Brownfield / greenfield", 13), ("P&ID", 10), ("Air P&ID", 11),
     ("Sub-system", 48), ("Comments", 40)]
add_type("cv", "Control / actuated valve", "2000-F00-LST-PP-10003", "", F)
for r in rows("2000-F00-LST-PP-10003", "Control Valve List", 1):
    if clean(r[4]): put("cv", F, r)

# ---- Line list (service codes expanded from the list's own lookup sheet) ----
svc = {}
for r in rows("2000-F00-LST-PP-10001", "List Inputs - Do Not Print", 1):
    if r and clean(r[0]) and clean(r[1]): svc.setdefault(clean(r[0]), clean(r[1]))
F = [("Line number", 0), ("Service", 3), ("Service description", -1), ("Size (DN)", 2), ("Pipe spec", 5), ("From", 10), ("To", 12),
     ("Area", 1), ("Insulation", 6), ("Insulation thickness (mm)", 7), ("Design pressure (kPag)", 17), ("Design temperature (°C)", 18),
     ("Operating pressure (kPag)", 42), ("Operating temperature (°C)", 43), ("Above / below ground", 19), ("NDE", 20), ("Pressure piping", 23),
     ("Pressure test", 24), ("Test pressure (kPag)", 34), ("Test medium", 35), ("Test duration", 36), ("Design standard", 33),
     ("Coating system", 28), ("Coating colour", 29), ("Corrosion allowance", 41), ("Pipe length (m)", 32), ("Stress critical", 45),
     ("Stage", 15), ("Brownfield / greenfield", 16), ("P&ID", 14), ("Sub-system", 46), ("Comments", 30)]
add_type("line", "Pipe line", "2000-F00-LST-PP-10001", "Rev 2", F)
for r in rows("2000-F00-LST-PP-10001", "Line List", 5):
    if not clean(r[0]) or not re.match(r"^\d", clean(r[0])): continue
    r = list(r) + [svc.get(clean(r[3]), "")]
    rec = [clean(r[c]) if c >= 0 else clean(r[-1]) for _, c in F]
    for i, (auto, man) in ((5, (11, 10)), (6, (13, 12))):   # From/To: "Auto" holds the full line/equipment tag,
        rec[i] = clean(r[auto]) or clean(r[man])             # "Manual" only a unique number or free text
    while rec and rec[-1] == "": rec.pop()
    DATA["line"].append(rec)

# ---- Manual valves ----
F = [("Valve tag", 7), ("Valve type", 1), ("Spec", 2), ("Size (DN)", 3), ("Code", 4), ("Area", 5), ("P&ID", 8), ("Brownfield / greenfield", 9),
     ("Stage", 10), ("Package", 11), ("Line number", 13), ("Revision", 0)]
add_type("mv", "Manual valve", "2000-F00-LST-PP-10002", "Rev 9", F)
for r in rows("2000-F00-LST-PP-10002", "Manual Valve List Tags", 1):
    if clean(r[7]): put("mv", F, r)
# manual valve models by code, from the summary sheet
models = {}
for r in rows("2000-F00-LST-PP-10002", "Manual Valve List", 7):
    code = clean(r[4]) if len(r) > 4 else ""
    if code: models[code] = [clean(r[i]) if i < len(r) else "" for i in (15, 16, 17, 18, 19)]
TYPES["mv"]["f"] += ["Supplier", "Manufacturer", "Model", "Model number", "General arrangement"]
for rec in DATA["mv"]:
    m = models.get(rec[4] if len(rec) > 4 else "")
    if m:
        rec += [""] * (12 - len(rec)); rec += m
        while rec and rec[-1] == "": rec.pop()

# ---- SPI (pipe specials), with the P&IDs each appears on ----
ws = openpyxl.load_workbook(find("2000-F00-LST-PP-10005"), read_only=True, data_only=True)["Full Pipe Specials List"]
allr = list(ws.iter_rows(values_only=True))
pids = {i: clean(v) for i, v in enumerate(allr[1]) if clean(v).startswith("2000-")}
F = [("SPI tag", 3), ("Type", 1), ("Description", 4), ("Service", 5), ("Pipe spec", 6), ("Size (DN)", 7), ("Stage 1 qty", 8), ("Stage 2 qty", 9),
     ("Total qty", 12), ("Make / model", 13), ("Specification", 14), ("Dimensions", 15), ("Inlet connection", 16), ("Discharge connection", 17),
     ("Design temperature (°C)", 18), ("Design pressure (kPag)", 19), ("Package", 20), ("GA drawing", 22), ("Comments", 23)]
add_type("spi", "Pipe special (SPI)", "2000-F00-LST-PP-10005", "Rev 1", F)
TYPES["spi"]["f"].append("P&IDs")
for r in allr[7:]:
    if not clean(r[3]).startswith("SP-"): continue
    rec = [clean(r[c]) if c < len(r) else "" for _, c in F]
    on = [p for i, p in pids.items() if i < len(r) and clean(r[i]) not in ("", "0")]
    rec += [""] * (len(F) - len(rec)); rec.append(", ".join(on))
    while rec and rec[-1] == "": rec.pop()
    DATA["spi"].append(rec)

# ---- Slurry hoses (Stage 1 and Stage 2 lists) ----
F1 = [("Hose tag", 3), ("Description", 4), ("Service", 8), ("Pipe spec", 9), ("Size (DN)", 10), ("Internal diameter (mm)", 11), ("Length (m)", 12),
      ("Location ref", 13), ("Specification", 14), ("Inlet connection", 15), ("Discharge connection", 16), ("Max bend radius", 17),
      ("Design temperature (°C)", 26), ("Design pressure (kPag)", 27), ("Package", 29), ("GA drawing", 30), ("Quantity", 7), ("Stage", -1)]
add_type("hose", "Slurry hose", "2000-F00-LST-PP-10032 / 12032", "", F1)
for doc, st, shift in (("2000-F00-LST-PP-10032", "Stage 1", 0), ("2000-F00-LST-PP-12032", "Stage 2", -1)):
    for r in rows(doc, "Slurry Hose List", 6):
        if not clean(r[3]).startswith("SP-"): continue
        rec = []
        for name, c in F1:
            if c < 0: rec.append(st); continue
            cc = c + (shift if c >= 18 else 0)       # Stage 2 list has no bend radius column
            rec.append(clean(r[cc]) if cc < len(r) else "")
        DATA["hose"].append(rec)

out = {"built": "tools/build_search.py", "types": TYPES, "data": DATA}
json.dump(out, open(OUT, "w"), ensure_ascii=False, separators=(",", ":"))
print({k: len(v) for k, v in DATA.items()}, round(os.path.getsize(OUT) / 1e6, 2), "MB")
