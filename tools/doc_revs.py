"""Reads the revision block (latest revision and its date) of every source document and writes tools/doc_revs.json.
Used by tools/build_issues.py to decide which of two clashing documents is newer (newer is taken as true).

Usage: python3 tools/doc_revs.py <folder holding the source .xlsx and .pdf files>
Where each date comes from:
  spreadsheets  the "Cover page" sheet, rows under DATE / REVISION / STATUS (latest date wins)
  MEL           no cover page, so the "Rev 0, Issued for Use, 26/12/2025" text at the top of the MEL sheet
  drawings      the title block revision rows, e.g. "1 | 04SEP26 | SP1 AS BUILT" (most common latest row across the sheets)
  reports       the revision table on the first pages, e.g. "10-Sept-24" (latest date), revision from the file name
Only document numbers, revisions, dates and status are kept (no names).
"""
import datetime, glob, json, os, re, sys, warnings
from collections import Counter
warnings.filterwarnings("ignore")
SRC = sys.argv[1]
OUT = os.path.join(os.path.dirname(__file__), "doc_revs.json")
MON = {m: i + 1 for i, m in enumerate("JAN FEB MAR APR MAY JUN JUL AUG SEP OCT NOV DEC".split())}

def find(doc, ext):
    m = [f for f in glob.glob(os.path.join(SRC, "*" + ext)) if doc in os.path.basename(f)]
    return m[0] if m else None

def cover(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[next(s for s in wb.sheetnames if "cover" in s.lower())]
    best = None
    for r in ws.iter_rows(values_only=True):
        v = [x for x in r if x not in (None, "")]
        if len(v) >= 3 and isinstance(v[0], datetime.datetime):
            row = (v[0].date(), str(v[1]), str(v[2]))
            if not best or row[0] > best[0] or (row[0] == best[0] and row[1] > best[1]): best = row
    return best

def mel(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["MEL"]
    for r in ws.iter_rows(max_row=4, values_only=True):
        for x in r:
            m = re.search(r"Rev\s*(\w+),\s*([^,]+),\s*(\d{1,2})/(\d{1,2})/(\d{4})", str(x or ""))
            if m: return (datetime.date(int(m[5]), int(m[4]), int(m[3])), m[1], m[2].strip())

def drawing(path):
    import pymupdf
    rows = Counter()
    for p in pymupdf.open(path):
        t = p.get_text().replace("\n", "|")
        revs = re.findall(r"\|\s*([0-9A-Z])\|\s*(\d{2})([A-Z]{3})(\d{2})\s*\|?\s*([A-Z0-9 ]+?)\|", t)
        if revs: rows[max(revs, key=lambda r: r[0])] += 1
    (rev, d, mo, y, st), n = rows.most_common(1)[0]
    return (datetime.date(2000 + int(y), MON[mo], int(d)), rev, st.strip())

def report(path):
    import pymupdf
    doc = pymupdf.open(path); t = " ".join(doc[i].get_text() for i in range(min(3, len(doc))))
    ds = []
    for d, mo, y in re.findall(r"\b(\d{1,2})-([A-Z][a-z]{2})[a-z]*-(\d{2})\b", t):
        if mo.upper() in MON: ds.append(datetime.date(2000 + int(y), MON[mo.upper()], int(d)))
    m = re.search(r"_(\w+?)_A", os.path.basename(path))
    return (max(ds), m[1] if m else "", "") if ds else None

DOCS = [  # key, title, document number, file extension, reader
    ("mel", "MEL", "2000-F00-LST-ME-10001", ".xlsx", mel),
    ("ins", "Instrument List", "2000-F00-LST-IC-11011", ".xlsx", cover),
    ("cv", "Control Valve List", "2000-F00-LST-PP-10003", ".xlsx", cover),
    ("line", "Line List", "2000-F00-LST-PP-10001", ".xlsx", cover),
    ("mv", "Manual Valve List", "2000-F00-LST-PP-10002", ".xlsx", cover),
    ("spi", "SPI List", "2000-F00-LST-PP-10005", ".xlsx", cover),
    ("hose1", "Slurry Hose List Stage 1", "2000-F00-LST-PP-10032", ".xlsx", cover),
    ("hose2", "Slurry Hose List Stage 2", "2000-F00-LST-PP-12032", ".xlsx", cover),
    ("pdc", "PDC", "2000-F00-DCR-PR-10002", ".pdf", report),
    ("cp1", "Control Philosophy Stage 1", "2000-F00-REP-PR-10001", ".pdf", report),
    ("cp2", "Control Philosophy Stage 2", "2000-F00-REP-PR-10002", ".pdf", report),
    ("pfd", "Stage 1 PFDs", "PFDs_combined", ".pdf", drawing),
]
out = {}
for key, title, num, ext, fn in DOCS:
    f = find(num, ext)
    r = fn(f) if f else None
    if not r: print("no revision block found for", title); continue
    out[key] = {"title": title, "number": num if "-" in num else "2000-F00-PFD-PR series", "rev": r[1], "date": r[0].isoformat(), "status": r[2]}
    print(f"{title:28} Rev {r[1]:3} {r[0]}  {r[2]}")
json.dump(out, open(OUT, "w"), indent=1)
